import os
import datetime
import openpyxl

def generate_fixtures():
    fixtures_dir = os.path.join(os.path.dirname(__file__), "fixtures")
    os.makedirs(fixtures_dir, exist_ok=True)

    # 1. sample_types.xlsx
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "Sheet1"
    headers = ["Name", "Age", "Score", "Active", "Created"]
    ws1.append(headers)
    ws1.append(["Alice", 30, 95.5, True, datetime.datetime(2023, 1, 15, 10, 30, 0)])
    ws1.append(["Bob", 25, 88.0, False, datetime.datetime(2023, 2, 20, 14, 0, 0)])
    ws1.append(["Charlie", 40, 72.25, True, datetime.datetime(2023, 3, 5, 9, 15, 0)])
    ws1.append(["Total", "=SUM(B2:B4)", "=AVERAGE(C2:C4)", True, None])
    wb1.save(os.path.join(fixtures_dir, "sample_types.xlsx"))

    # 2. multilingual.xlsx
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Languages"
    ws2.append(["Language", "Sample Text"])
    ws2.append(["English", "Hello World & Friends <tag> 'quoted' \"double\""])
    ws2.append(["Chinese", "你好世界，很高兴认识你"])
    ws2.append(["Japanese", "こんにちは世界、よろしく"])
    ws2.append(["Arabic", "مرحبا بالعالم"])
    ws2.append(["Russian", "Привет мир, добро пожаловать"])
    ws2.append(["German", "Grüße aus München & Köln äöüß"])
    ws2.append(["Emoji", "🚀 🦀 🐍 ⚡ 📊"])
    wb2.save(os.path.join(fixtures_dir, "multilingual.xlsx"))

    # 3. multi_sheet.xlsx
    wb3 = openpyxl.Workbook()
    ws_summary = wb3.active
    ws_summary.title = "Summary"
    ws_summary.append(["Metric", "Value"])
    ws_summary.append(["Total Users", 1000])

    ws_2023 = wb3.create_sheet(title="Data 2023")
    ws_2023.append(["Year", "Sales"])
    ws_2023.append([2023, 500000])

    ws_2024 = wb3.create_sheet(title="Data 2024")
    ws_2024.append(["Year", "Sales"])
    ws_2024.append([2024, 750000])

    wb3.active = ws_2023  # Set middle sheet as active
    wb3.save(os.path.join(fixtures_dir, "multi_sheet.xlsx"))

    # 4. sparse.xlsx
    wb4 = openpyxl.Workbook()
    ws4 = wb4.active
    ws4.title = "SparseSheet"
    ws4["A1"] = 100
    ws4["C1"] = 300
    ws4["A5"] = 500
    ws4["D10"] = "Far Away"
    wb4.save(os.path.join(fixtures_dir, "sparse.xlsx"))

    # 5. corrupted.xlsx (plain text file, invalid zip)
    with open(os.path.join(fixtures_dir, "corrupted.xlsx"), "w") as f:
        f.write("This is definitely not a valid XLSX or ZIP file.")

    print(f"Successfully generated fixtures in {fixtures_dir}")

if __name__ == "__main__":
    generate_fixtures()
